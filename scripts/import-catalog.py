#!/usr/bin/env python3
"""Import supplier snapshots. Only the sanitised catalogue belongs under public/."""
from __future__ import annotations

import argparse
from collections import Counter
import csv
from datetime import date
from decimal import Decimal, InvalidOperation
import hashlib
import io
import json
from pathlib import Path
import re
import sys
import unicodedata


BRANDS = ["Logitech", "Kingston", "TP-Link", "D-Link", "Genius", "Trust", "Samsung", "Motorola", "Xiaomi", "Apple", "iPhone", "Intel", "AMD", "Asus", "ASRock", "Gigabyte", "Acer", "Lenovo", "HP", "Dell", "Seagate", "Western Digital", "WD", "Adata", "Hiksemi", "Dahua", "Hikvision", "Epson", "Brother", "Pantum", "Yaxun", "Sunshine", "Philips", "ViewSonic", "Cooler Master", "Thermaltake", "Redragon", "XPG", "Patriot", "Crucial", "Sandisk", "Noga", "Nisuta", "Delta", "Booru", "BYH", "Netac", "Razer", "NZXT", "Toshiba", "Solarmax", "Zowie"]


def clean(value):
    return re.sub(r"\s+", " ", str(value or "")).strip()


def folded(value):
    return "".join(c for c in unicodedata.normalize("NFD", clean(value).upper()) if not unicodedata.combining(c))


def number(value):
    if value is None or value == "":
        return None
    return float(Decimal(str(value)))


def arg_money(value):
    return float(Decimal(clean(value).replace("$", "").replace(".", "").replace(",", ".").strip()))


def brand_for(name):
    name = folded(name).replace("TP LINK", "TP-LINK")
    for brand in BRANDS:
        if re.search(r"(?<![A-Z0-9])" + re.escape(folded(brand)) + r"(?![A-Z0-9])", name):
            return brand
    return None


def category_for(source, name):
    text = folded(source)
    item = folded(name)
    if "OUTLET" in text:
        return "Outlet"
    exact_categories = {
        "PLACAS DE VIDEO": "Componentes de PC",
        "PLACA DE VIDEO": "Componentes de PC",
        "ACCESS POINT": "Redes y conectividad",
        "MODEMS": "Redes y conectividad",
        "PANTALLAS CELULARES": "Repuestos de celulares",
        "PANTALLAS SMARTWATCH": "Repuestos y accesorios",
        "JACK DE CARGA DC JACK": "Repuestos y accesorios",
        "BATERIAS GEL": "Repuestos y accesorios",
        "BATERIAS INALAMBRICOS": "Repuestos y accesorios",
        "BATERIAS MP3 MP4 IPOD": "Repuestos y accesorios",
        "ARDUINO MODULOS": "Componentes electrónicos",
        "CONSOLAS": "Consolas y accesorios",
        "CONSOLAS Y ACCESORIOS": "Consolas y accesorios",
    }
    if text in exact_categories:
        return exact_categories[text]
    if text in {"NOTEBOOKS Y PORTATILES", "COMPUTADORAS", "NETBOOKS NOTEBOOKS", "PCS ARMADAS"}:
        return "Notebooks y computadoras"
    if any(x in text for x in ("NOTEBOOK", "NETBOOK")):
        return "Accesorios y repuestos de notebooks"
    if "CONSOL" in text:
        if any(x in text for x in ("FLEX", "BATERIA", "COOLER", "MOTOR", "MEMORIA", "REPUESTO", "MODULO", "PANTALLA")):
            return "Repuestos de consolas"
        return "Consolas y accesorios"
    if "TABLET" in text and any(x in text for x in ("PANTALLA", "MODULO", "BATERIA", "FLEX", "TACTIL", "REPUESTO")):
        return "Repuestos de tablets"
    if not text:
        is_part = any(x in item for x in ("PANTALLA", "MODULO", "BATERIA", "FLEX", "TACTIL", "REPUESTO", "BISAGRA", "JACK DE CARGA"))
        if is_part and any(x in item for x in ("NOTEBOOK", "NETBOOK")):
            return "Accesorios y repuestos de notebooks"
        if is_part and "TABLET" in item:
            return "Repuestos de tablets"
        if is_part and any(x in item for x in ("CONSOLA", "PLAYSTATION", "XBOX", "NINTENDO")):
            return "Repuestos de consolas"
        if "ACCESS POINT" in item:
            return "Redes y conectividad"
        if any(x in item for x in ("PLACA DE VIDEO", "PLACAS DE VIDEO")):
            return "Componentes de PC"
    if text in {"MONITORES", "TELEVISORES", "PROYECTORES"}:
        return "Monitores y pantallas"
    if "TABLET" in text and any(x in text for x in ("FUNDA", "SOPORTE", "FILM", "VIDRIO", "CARGADOR")):
        return "Accesorios para celulares y tablets"
    if any(x in text for x in ("FUNDAS DE CELULAR", "FILMS PROTECTORES", "VIDRIOS CELULARES", "CARGADORES CELULARES", "SOPORTES PARA CELULAR", "CABLES CARGADORES")):
        return "Accesorios para celulares y tablets"
    if "REPUESTOS" in text and any(x in text for x in ("JOYSTICK", "TECLADOS PC")):
        return "Repuestos y accesorios de periféricos"
    if any(x in text for x in ("CARGADORES CAMARAS", "BATERIAS CAMARAS", "FUNDAS DE CAMARAS", "TRANSFORMADOR", "SOPORTES PARA TV", "SOPORTES PARA PROYECTOR")):
        return "Accesorios"
    if any(x in text for x in ("HDD", "SSD", "MEMORIA", "PENDRIVE", "DISCO", "ALMACENAMIENTO")):
        return "Almacenamiento y memorias"
    if any(x in text for x in ("CONECTIVIDAD", "RED", "ROUTER", "SWITCH", "WIFI", "ANTENA", "ACCESS POINT")):
        return "Redes y conectividad"
    if any(x in text for x in ("MOUSE", "TECLADO", "WEBCAM", "JOYSTICK", "PERIFERICO")):
        return "Periféricos"
    if any(x in text for x in ("AURICULAR", "PARLANTE", "MICROFONO", "AUDIO")):
        return "Audio"
    if any(x in text for x in ("MOTHER", "PROCESADOR", "PLACA DE VIDEO", "PLACAS DE VIDEO", "GABINETE", "COOLER", "FUENTE", "CASE")):
        return "Componentes de PC"
    if any(x in text for x in ("IMPRES", "CONSUMIBLE", "CARTUCHO", "TONER", "TINTAS", "CINTAS MATRICIALES")):
        return "Impresión"
    if any(x in text for x in ("CAMARA", "SEGURIDAD", "CASA INTELIGENTE", "HOGAR", "ILUMINACION")):
        return "Hogar y seguridad"
    if any(x in text for x in ("HERRAMIENTA", "TECNICO", "INSUMO", "DESTORNILLADOR", "SOLDAD", "REBALLING", "GRASA SILICONADA", "LIMPIADOR")):
        return "Herramientas e insumos"
    if any(x in text for x in ("REPUESTO", "MODULO", "FLEX", "BATERIA", "PIN DE CARGA", "TAPA", "BUZZER", "ZOCALO", "PLACA DE CARGA", "JACK", "TACTIL")):
        return "Repuestos de celulares"
    if "CELULAR" in text or "TABLET" in text:
        return "Celulares y tablets"
    if any(x in item for x in ("DESTORNILLADOR", "ALICATE", "YAXUN", "SUNSHINE", "SOLDAD", "FLUX", "ISOPROPIL")):
        return "Herramientas e insumos"
    if any(x in item for x in ("MODULO", "FLEX", "BATERIA", "PIN DE CARGA", "TAPA ", "BUZZER", "ZOCALO", "PLACA DE CARGA", "JACK", "TACTIL")):
        return "Repuestos de celulares"
    if any(x in item for x in ("TECLADO", "MOUSE", "WEBCAM")):
        return "Periféricos"
    if any(x in item for x in ("AURICULAR", "PARLANTE", "MICROFONO")):
        return "Audio"
    return "Accesorios"


def display_name(name):
    # Product names and model variants remain verbatim. Altering case can damage model IDs.
    return clean(name)


def source_record(path, source_id, source_date):
    return {"id": source_id, "fileName": path.name, "sourceDate": source_date, "sha256": hashlib.sha256(path.read_bytes()).hexdigest()}


def nb_source_date(path):
    match = re.search(r"(\d{2})_(\d{2})_(\d{4})", path.name)
    return f"{match[3]}-{match[2]}-{match[1]}" if match else None


def identifier(value):
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if value != int(value):
            raise ValueError(f"Identificador numérico no entero: {value}")
        return str(int(value))
    return clean(value)


def read_nb_rows(path, xls_deps=None):
    """Detect real BIFF XLS or zipped XLSX from file contents, not its extension."""
    with path.open("rb") as handle:
        signature = handle.read(8)
    notes = []
    if signature == bytes.fromhex("d0cf11e0a1b11ae1"):
        if xls_deps is not None and xls_deps.is_dir():
            sys.path.insert(0, str(xls_deps.resolve()))
        try:
            import xlrd
            from xlrd.compdoc import CompDocError
        except ModuleNotFoundError as error:
            raise RuntimeError("Para XLS real se requiere xlrd. Instalalo en un directorio aislado y pasá --xls-deps RUTA; no es necesario modificar el runtime.") from error
        diagnostic = io.StringIO()
        try:
            book = xlrd.open_workbook(str(path), logfile=diagnostic)
        except CompDocError as error:
            # The NB BIFF export reuses OLE sectors. Retry this known container
            # defect only; row schema, IDs and every cost are still validated below.
            if not re.fullmatch(r"Workbook corruption: seen\[\d+\] == \d+", str(error)):
                raise
            notes.append(str(error) + "; lectura tolerante del contenedor OLE, con validación posterior de esquema, IDs y costos.")
            book = xlrd.open_workbook(str(path), logfile=diagnostic, ignore_workbook_corruption=True)
        try:
            sheets = [sheet for sheet in book.sheets() if sheet.nrows]
            if len(sheets) != 1:
                raise ValueError("NB: se esperaba una única hoja con datos.")
            sheet = sheets[0]
            rows = [tuple(None if value == "" else value for value in sheet.row_values(index)) for index in range(sheet.nrows)]
            metadata = {"format": "xls-biff", "sheetName": sheet.name, "readerNotes": notes}
        finally:
            book.release_resources()
    elif signature.startswith(b"PK"):
        from openpyxl import load_workbook
        with path.open("rb") as handle:
            book = load_workbook(handle, data_only=True, read_only=True)
            sheet = book.active
            rows = list(sheet.iter_rows(values_only=True))
            metadata = {"format": "xlsx", "sheetName": sheet.title, "readerNotes": notes}
            book.close()
    else:
        raise ValueError("NB: el archivo no es XLS BIFF ni XLSX; no se interpretará por extensión.")
    return rows, metadata


def nb_import(path, xls_deps=None):
    rows, metadata = read_nb_rows(path, xls_deps)
    if not any(row and folded(row[0]) == "CODIGO" for row in rows[:10]):
        raise ValueError("NB: esta lista no tiene el esquema detallado CODIGO/MONEDA/PRECIO FINAL. Para la exportación Identificador de URL/Precio/Costo usá --update-nb junto al catálogo privado anterior.")
    records = []
    category = "Accesorios"
    for row_number, row in enumerate(rows, start=1):
        if len(row) < 9:
            continue
        if row[0] is None and row[1] is not None:
            category = clean(row[1])
        if not isinstance(row[0], (int, float)) or not row[2]:
            continue
        code = identifier(row[0])
        currency = {"U$S": "USD", "USD": "USD", "$": "ARS", "ARS": "ARS"}.get(clean(row[6]))
        if currency is None:
            raise ValueError(f"NB row {row_number}: unknown currency {row[6]!r}")
        vat_rate = number(clean(row[3]).replace("%", ""))
        vat_rate = vat_rate / 100 if vat_rate is not None else None
        gross = number(row[8])
        if gross is None or gross <= 0:
            raise ValueError(f"NB row {row_number}: missing/invalid PRECIO FINAL")
        records.append({
            "id": f"nb-{code}", "supplier": "nb", "supplierCode": code,
            "sourceRow": row_number, "sourcePage": None,
            "name": clean(row[2]), "manufacturerId": identifier(row[1]) or None,
            "category": category_for(category, row[2]), "sourceCategory": category,
            "brand": brand_for(row[2]), "image": clean(row[12]) or None,
            "gallery": [url.strip() for url in clean(row[21]).split(",") if url.strip()] if len(row) > 21 else [],
            "warranty": clean(row[5]) or None,
            "ean": clean(row[14]) or None, "upc": clean(row[15]) or None,
            "cost": {"currency": currency, "net": number(row[7]), "gross": gross, "vatRate": vat_rate, "vatIncluded": True, "internalTaxFromSource": number(row[20])},
            "quantityPrices": [{"minQuantity": 1, "gross": gross}],
            "supplierStockSnapshot": number(row[4]),
            "supplierStockNote": "Dato histórico del proveedor. No representa stock propio ni disponibilidad actual.",
        })
    source = source_record(path, "nb", nb_source_date(path))
    source.update(metadata)
    return records, source


def parse_nb_cost_rows(rows, declared_currency=None):
    """Read NB's store export; Costo is the wholesale final cost, never Precio."""
    if not rows:
        raise ValueError("NB: archivo sin filas.")
    headers = [folded(value) for value in rows[0]]
    required = ["IDENTIFICADOR DE URL", "NOMBRE", "SKU", "PRECIO", "COSTO"]
    if any(name not in headers for name in required) or len(headers) != len(set(headers)):
        raise ValueError("NB: encabezado de exportación no reconocido o duplicado.")
    columns = {name: index for index, name in enumerate(headers)}
    if declared_currency not in {"ARS", "USD"} and "MONEDA" not in columns:
        raise ValueError("La exportación NB no declara moneda. Indicá --nb-currency ARS sólo para una lista confirmada en pesos.")
    updates = []
    seen = set()
    for row_number, row in enumerate(rows[1:], start=2):
        if not any(value is not None and value != "" for value in row):
            continue
        if len(row) < len(headers):
            raise ValueError(f"NB fila {row_number}: faltan columnas.")
        code = identifier(row[columns["IDENTIFICADOR DE URL"]])
        if not re.fullmatch(r"\d+", code):
            raise ValueError(f"NB fila {row_number}: Identificador de URL inválido.")
        product_id = "nb-" + code
        if product_id in seen:
            raise ValueError(f"NB fila {row_number}: ID repetido {product_id}.")
        seen.add(product_id)
        raw_currency = clean(row[columns["MONEDA"]]) if "MONEDA" in columns else None
        currency = {"U$S": "USD", "USD": "USD", "$": "ARS", "ARS": "ARS"}.get(raw_currency) if raw_currency else declared_currency
        if currency not in {"ARS", "USD"} or (raw_currency and declared_currency and currency != declared_currency):
            raise ValueError(f"NB fila {row_number}: moneda desconocida o contradictoria.")
        try:
            cost = Decimal(str(row[columns["COSTO"]]))
        except (InvalidOperation, TypeError, ValueError) as error:
            raise ValueError(f"NB fila {row_number}: Costo inválido.") from error
        if not cost.is_finite() or cost <= 0:
            raise ValueError(f"NB fila {row_number}: Costo debe ser positivo y finito.")
        name = clean(row[columns["NOMBRE"]])
        if not name:
            raise ValueError(f"NB fila {row_number}: falta el nombre.")
        updates.append({
            "id": product_id, "name": name, "manufacturerId": identifier(row[columns["SKU"]]),
            "sourceRow": row_number, "currency": currency, "gross": float(cost),
            "internalTaxFromSource": number(row[columns["IMPUESTO INTERNO"]]) if "IMPUESTO INTERNO" in columns else None,
        })
    if not updates:
        raise ValueError("NB: no se encontraron productos con costos válidos.")
    return updates


def merge_nb_costs(catalog, updates, source, as_of):
    """Update known NB IDs in place; preserve every public/product metadata field."""
    indexed = {record["id"]: record for record in catalog["products"] if record["supplier"] == "nb"}
    if len(indexed) != sum(record["supplier"] == "nb" for record in catalog["products"]):
        raise ValueError("El catálogo anterior tiene IDs NB repetidos.")
    incoming = {update["id"] for update in updates}
    if len(incoming) != len(updates):
        raise ValueError("La actualización contiene IDs NB repetidos.")
    unknown = sorted(incoming - indexed.keys())
    if unknown:
        raise ValueError("La lista incluye IDs NB sin metadatos previos; importá primero su ficha detallada: " + ", ".join(unknown))
    for update in updates:
        record = indexed[update["id"]]
        if clean(record["name"]) != update["name"] or (record.get("manufacturerId") and record["manufacturerId"] != update["manufacturerId"]):
            raise ValueError(f"La variante de {update['id']} no coincide con la ficha previa; no se actualizó el catálogo.")
    other_before = json.dumps([record for record in catalog["products"] if record["supplier"] != "nb"], sort_keys=True, ensure_ascii=False)
    for update in updates:
        record = indexed[update["id"]]
        old_cost = record["cost"]
        if "costFromSource" not in record:
            record["costFromSource"] = dict(old_cost)
            record["costFromSourceSourceId"] = record.get("costSource", {}).get("sourceId", "nb")
        if "quantityPricesFromSource" not in record:
            record["quantityPricesFromSource"] = record["quantityPrices"]
        record["cost"] = {
            "currency": update["currency"], "net": None, "gross": update["gross"],
            "vatRate": old_cost.get("vatRate"), "vatIncluded": old_cost.get("vatIncluded"),
            "internalTaxFromSource": update["internalTaxFromSource"],
        }
        record["quantityPrices"] = [{"minQuantity": 1, "gross": update["gross"]}]
        record["costSource"] = {
            "sourceId": source["id"], "sourceRow": update["sourceRow"], "column": "Costo",
            "currencyFrom": "MONEDA" if source.get("hasCurrencyColumn") else "CLI --nb-currency, moneda confirmada por el propietario",
            "vatTreatmentFrom": "Ficha NB detallada anterior del mismo SKU; no se agrega nuevamente IVA al costo final.",
        }
        record["costUpdatedAt"] = as_of
        record["missingFromLatestNbList"] = False
        record["costRequiresConfirmation"] = old_cost.get("vatIncluded") is not True
        record["latestNbListSourceId"] = source["id"]
        record.pop("costOverrideNote", None)
    missing = sorted(indexed.keys() - incoming)
    for product_id in missing:
        indexed[product_id]["missingFromLatestNbList"] = True
        indexed[product_id]["costRequiresConfirmation"] = True
        indexed[product_id]["latestNbListSourceId"] = source["id"]
        indexed[product_id].setdefault("costSource", {"sourceId": "nb", "sourceRow": indexed[product_id].get("sourceRow"), "column": "PRECIO FINAL", "note": "Costo de la lista NB anterior; el producto no figura en la última lista y requiere confirmación."})
    other_after = json.dumps([record for record in catalog["products"] if record["supplier"] != "nb"], sort_keys=True, ensure_ascii=False)
    if other_before != other_after:
        raise AssertionError("Se alteraron registros de otros proveedores.")
    catalog["sources"] = [previous for previous in catalog["sources"] if previous["id"] != source["id"]] + [source]
    audit = {
        "updatedAt": as_of, "updatedCount": len(updates), "missingCount": len(missing), "missingIds": missing,
        "unchangedOtherSuppliers": sum(record["supplier"] != "nb" for record in catalog["products"]),
        "otherSuppliersSha256": hashlib.sha256(other_after.encode("utf8")).hexdigest(),
        "costColumn": "Costo", "retailPricesPublished": False,
    }
    catalog["nbUpdateAudit"] = audit
    return audit


def update_nb_catalog(private_path, path, declared_currency, xls_deps, as_of):
    if not private_path.is_file():
        raise ValueError("--update-nb requiere el catálogo privado previo para conservar fotos, variantes y metadatos.")
    catalog = json.loads(private_path.read_text(encoding="utf-8"))
    rows, metadata = read_nb_rows(path, xls_deps)
    updates = parse_nb_cost_rows(rows, declared_currency)
    source = source_record(path, "nb-costs-" + hashlib.sha256(path.read_bytes()).hexdigest()[:12], nb_source_date(path))
    source.update(metadata)
    source.update({"costColumn": "Costo", "currencyDeclaration": declared_currency, "hasCurrencyColumn": "MONEDA" in [folded(value) for value in rows[0]], "productCount": len(updates)})
    audit = merge_nb_costs(catalog, updates, source, as_of)
    write_json(private_path, catalog)
    return audit


def distriland_import(path):
    numbered = [(index + 1, clean(line)) for index, line in enumerate(path.read_text(encoding="utf-8-sig").splitlines()) if clean(line)]
    lines = [line for _, line in numbered]
    records = []
    used = set()
    for index, line in enumerate(lines):
        if not re.fullmatch(r"\$\s*\d[\d.]*?(?:,\d{1,2})?", line) or index == 0:
            continue
        # The pasted store listing repeats the image alt/title then the visible name.
        # The closest preceding non-price line is the visible, exact product variant.
        name = re.sub(r" - comprar online$", "", lines[index - 1], flags=re.I)
        if name in {"Detalles", "Comprar"} or len(name) < 5:
            raise ValueError(f"DistriLand: unable to identify product before {line}")
        identifier = hashlib.sha256(name.encode("utf-8")).hexdigest()[:14]
        product_id = f"dl-{identifier}"
        if product_id in used:
            raise ValueError(f"DistriLand: repeated product variant {name}")
        used.add(product_id)
        records.append({
            "id": product_id, "supplier": "distriland", "supplierCode": None,
            "sourceRow": numbered[index][0], "sourcePage": None,
            "name": name, "manufacturerId": None,
            "category": category_for("", name), "sourceCategory": None,
            "brand": brand_for(name), "image": None, "gallery": [],
            "cost": {"currency": "ARS", "net": None, "gross": arg_money(line), "vatRate": None, "vatIncluded": None},
            "quantityPrices": [{"minQuantity": 1, "gross": arg_money(line)}],
            "supplierStockSnapshot": None,
            "notes": "Precio copiado de la web. No se acredita tarifa mayorista ni tratamiento del IVA. Confirmar antes de cotizar.",
        })
    return records, source_record(path, "distriland", None)


def retained_pdf_images(private_path, pdf_path, image_dir, image_prefix):
    """Keep already extracted photos only for the identical supplier PDF snapshot."""
    if not private_path.is_file():
        return {}
    previous = json.loads(private_path.read_text(encoding="utf-8"))
    prior_source = next((source for source in previous.get("sources", []) if source.get("id") == "tecnoprices"), None)
    if not prior_source or prior_source.get("sha256") != hashlib.sha256(pdf_path.read_bytes()).hexdigest():
        return {}
    prefix = image_prefix.rstrip("/") + "/"
    retained = {}
    for record in previous.get("products", []):
        image = record.get("image")
        if record.get("supplier") != "tecnoprices" or not image or not image.startswith(prefix):
            continue
        filename = image[len(prefix):]
        # Only files directly within the configured product-image directory qualify.
        if filename != Path(filename).name or not (image_dir / filename).is_file():
            continue
        retained[record["id"]] = image
    return retained


def tecnoprices_import(path, image_dir, image_prefix, extract_images=True, retained_images=None):
    import pdfplumber
    from pypdf import PdfReader

    records = []
    reader = PdfReader(path) if extract_images else None
    image_hashes = {}
    image_bytes = 0
    category = None
    source_date = None
    image_count = 0
    with pdfplumber.open(path) as pdf:
        for page_number, page in enumerate(pdf.pages, start=1):
            if abs(page.width - 841.89) > 2:
                raise ValueError(f"TecnoPrices page {page_number}: unsupported layout")
            words = page.extract_words(x_tolerance=1.5, y_tolerance=2)
            if page_number == 1:
                stamp = re.search(r"Generado el (\d{2})/(\d{2})/(\d{4}) (\d{2}:\d{2})", page.extract_text() or "")
                source_date = f"{stamp[3]}-{stamp[2]}-{stamp[1]}T{stamp[4]}:00-03:00" if stamp else None
            body = [word for word in words if 86 < word["top"] < page.height - 30]
            sku_words = [word for word in body if word["x0"] < 60 and re.fullmatch(r"\d+", word["text"])]
            # Category labels begin in the SKU column; product descriptions begin at x=100.
            heading_starts = [word for word in body if word["x0"] < 35 and not word["text"].isdigit()]
            headings = []
            for start in heading_starts:
                same_line = sorted([word for word in body if abs(word["top"] - start["top"]) < 1], key=lambda word: word["x0"])
                label = re.sub(r"\s*\(continuacion\)\s*$", "", clean(" ".join(word["text"] for word in same_line)), flags=re.I)
                headings.append((start["top"], label))
            for sku_word in sku_words:
                preceding = [heading for heading in headings if heading[0] < sku_word["top"]]
                if preceding:
                    category = preceding[-1][1]
                if not category:
                    raise ValueError(f"TecnoPrices page {page_number}: missing category")
                center = (sku_word["top"] + sku_word["bottom"]) / 2
                row_words = [word for word in body if abs((word["top"] + word["bottom"]) / 2 - center) < 12.5]
                description = sorted([word for word in row_words if 98 <= word["x0"] < 562], key=lambda word: (round(word["top"] / 3), word["x0"]))
                name = clean(" ".join(word["text"] for word in description))
                tax = [word for word in row_words if 562 <= word["x0"] < 594 and "%" in word["text"]]
                prices = []
                for minimum, left, right in [(1, 594, 650), (3, 650, 707), (5, 707, 763), (10, 763, 820)]:
                    price_words = [word for word in row_words if left <= (word["x0"] + word["x1"]) / 2 < right and re.fullmatch(r"\d[\d.]*,\d{2}", word["text"])]
                    if len(price_words) > 1:
                        raise ValueError(f"TecnoPrices page {page_number} SKU {sku_word['text']}: ambiguous tier x{minimum}")
                    if price_words:
                        prices.append({"minQuantity": minimum, "gross": arg_money(price_words[0]["text"])})
                if not name or len(tax) != 1 or not prices or prices[0]["minQuantity"] != 1:
                    raise ValueError(f"TecnoPrices page {page_number} SKU {sku_word['text']}: incomplete row")
                image_url = (retained_images or {}).get(f"tp-{sku_word['text']}")
                matching_images = [entry for entry in page.images if 60 < entry["x0"] < 98 and abs((entry["top"] + entry["bottom"]) / 2 - center) < 3]
                if extract_images and len(matching_images) == 1:
                    image = reader.pages[page_number - 1].images["/" + matching_images[0]["name"]]
                    digest = hashlib.sha256(image.data).hexdigest()
                    if digest not in image_hashes:
                        image_dir.mkdir(parents=True, exist_ok=True)
                        extension = Path(image.name).suffix.lower()
                        if extension not in {".jpg", ".jpeg", ".png"}:
                            raise ValueError(f"Unsupported embedded image {image.name}")
                        filename = f"tp-{sku_word['text']}{extension}"
                        (image_dir / filename).write_bytes(image.data)
                        image_hashes[digest] = image_prefix.rstrip("/") + "/" + filename
                        image_bytes += len(image.data)
                        image_count += 1
                    image_url = image_hashes[digest]
                records.append({
                    "id": f"tp-{sku_word['text']}", "supplier": "tecnoprices", "supplierCode": sku_word["text"],
                    "sourceRow": None, "sourcePage": page_number,
                    "name": name, "manufacturerId": None,
                    "category": category_for(category, name), "sourceCategory": category,
                    "brand": brand_for(name), "image": image_url, "gallery": [],
                    "cost": {"currency": "ARS", "net": None, "gross": prices[0]["gross"], "vatRate": arg_money(tax[0]["text"].replace("%", "")) / 100, "vatIncluded": True},
                    "quantityPrices": prices,
                    "supplierStockSnapshot": None,
                    "supplierStockNote": "El PDF incluye productos con stock al generarse. Requiere nueva confirmación antes de tomar un pedido.",
                })
            if page_number % 25 == 0:
                print(f"TecnoPrices: {page_number}/{len(pdf.pages)} páginas procesadas", file=sys.stderr, flush=True)
            page.close()
    return records, source_record(path, "tecnoprices", source_date), {"uniqueImages": image_count, "imageBytes": image_bytes, "retainedProductImages": len(retained_images or {})}


def public_product(record):
    description = f"Referencia {record['manufacturerId']}. " if record.get("manufacturerId") else ""
    description += "Producto a pedido. Confirmamos disponibilidad y cotización antes de la compra."
    return {
        "id": record["id"], "name": display_name(record["name"]), "category": record["category"],
        "brand": record["brand"], "image": record["image"], "description": description,
        "price": None, "currency": "ARS",
    }


def apply_nb_ars(records, path):
    """Use manually exported final ARS costs. Never infer an exchange rate."""
    indexed = {record["id"]: record for record in records if record["supplier"] == "nb"}
    seen = set()
    with path.open(encoding="utf-8-sig", newline="") as handle:
        rows = csv.DictReader(handle)
        if not {"id", "costARS"}.issubset(set(rows.fieldnames or [])):
            raise ValueError("El CSV NB ARS requiere columnas id,costARS.")
        for row in rows:
            product_id = clean(row["id"])
            if product_id not in indexed or product_id in seen:
                raise ValueError(f"ID NB inexistente o duplicado en CSV ARS: {product_id}")
            seen.add(product_id)
            gross = Decimal(clean(row["costARS"]))
            if not gross.is_finite() or gross <= 0:
                raise ValueError(f"Costo ARS inválido: {product_id}")
            record = indexed[product_id]
            record["costFromSource"] = record["cost"]
            record["cost"] = {"currency": "ARS", "net": None, "gross": float(gross), "vatRate": record["cost"]["vatRate"], "vatIncluded": True}
            record["quantityPricesFromSource"] = record["quantityPrices"]
            record["quantityPrices"] = [{"minQuantity": 1, "gross": float(gross)}]
            record["costOverrideNote"] = "Costo final ARS con IVA cargado manualmente desde CSV. Verificar vigencia antes de cotizar."
    return source_record(path, "nb-ars-overrides", None)


def initial_order(records):
    # A small, deliberate selection of ordinary technology products starts the shop.
    # Full supplier ordering is retained privately; all other variants remain searchable.
    priorities = ["nb-119109", "nb-123518", "nb-123594", "nb-122626", "nb-117544", "nb-118900", "nb-104939", "nb-121073", "nb-119170", "nb-102278", "nb-121091", "nb-121145", "nb-123697", "nb-118027", "nb-120470", "nb-118875"]
    chosen = [record for key in priorities for record in records if record["id"] == key]
    ids = {record["id"] for record in chosen}
    return chosen + [record for record in records if record["id"] not in ids]


def write_json(path, content):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(content, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    mode = parser.add_mutually_exclusive_group(required=True)
    mode.add_argument("--nb", type=Path, help="Lista NB detallada XLSX o XLS con CODIGO/MONEDA/PRECIO FINAL")
    mode.add_argument("--update-nb", type=Path, help="Actualizar sólo costos NB desde XLS/XLSX Identificador de URL/Precio/Costo, conservando fotos y catálogo público")
    parser.add_argument("--nb-currency", choices=["ARS", "USD"], help="Moneda confirmada para una exportación NB sin columna MONEDA; no convierte importes")
    parser.add_argument("--xls-deps", type=Path, default=Path(__file__).resolve().parents[1] / "artifacts/python-deps", help="Directorio aislado con xlrd, si no está instalado")
    parser.add_argument("--nb-ars", type=Path, help="CSV opcional id,costARS con costos finales NB en ARS, IVA incluido")
    parser.add_argument("--distriland", type=Path, help="Texto pegado DistriLand UTF-8")
    parser.add_argument("--tecnoprices", type=Path, help="Listado TecnoPrices PDF")
    parser.add_argument("--private-output", type=Path, default=Path(".private/catalog-costs.json"))
    parser.add_argument("--public-output", type=Path, default=Path("public/tienda/catalog.json"))
    parser.add_argument("--image-dir", type=Path, default=Path("public/tienda/assets/products"))
    parser.add_argument("--image-prefix", default="./assets/products")
    parser.add_argument("--skip-pdf-images", action="store_true")
    parser.add_argument("--as-of", default=date.today().isoformat(), help="Fecha de importación YYYY-MM-DD")
    args = parser.parse_args()
    date.fromisoformat(args.as_of)
    # Fail closed: the private wholesale-cost snapshot must never be a public asset.
    public_root = args.public_output.resolve().parents[1]
    if args.private_output.resolve().is_relative_to(public_root):
        raise ValueError("El catálogo privado debe guardarse fuera de public/.")
    if args.update_nb:
        if args.nb_ars or args.distriland or args.tecnoprices:
            parser.error("--update-nb actualiza sólo NB; no combines listas de otros proveedores ni --nb-ars.")
        audit = update_nb_catalog(args.private_output, args.update_nb, args.nb_currency, args.xls_deps, args.as_of)
        print(json.dumps({"mode": "update-nb", "privateOutput": str(args.private_output), "publicCatalogUnchanged": True, **audit}, ensure_ascii=False))
        return
    if not args.distriland or not args.tecnoprices:
        parser.error("La importación completa requiere --distriland y --tecnoprices; para actualizar sólo costos usá --update-nb.")
    if args.nb_currency:
        parser.error("--nb-currency corresponde al modo --update-nb; la lista NB detallada declara su moneda.")
    records, nb_source = nb_import(args.nb, args.xls_deps)
    dl_records, dl_source = distriland_import(args.distriland)
    retained_images = retained_pdf_images(args.private_output, args.tecnoprices, args.image_dir, args.image_prefix) if args.skip_pdf_images else {}
    tp_records, tp_source, image_stats = tecnoprices_import(args.tecnoprices, args.image_dir, args.image_prefix, not args.skip_pdf_images, retained_images)
    records.extend(dl_records)
    records.extend(tp_records)
    ids = [record["id"] for record in records]
    if len(ids) != len(set(ids)):
        raise ValueError("Hay IDs de producto repetidos; no se escribió ningún JSON.")
    sources = [nb_source, dl_source, tp_source]
    if args.nb_ars:
        sources.append(apply_nb_ars(records, args.nb_ars))
    write_json(args.private_output, {"version": 1, "importedAt": args.as_of, "sources": sources, "products": records})
    write_json(args.public_output, {"version": 1, "updatedAt": args.as_of, "products": [public_product(record) for record in initial_order(records)]})
    print(json.dumps({"products": len(records), "bySupplier": dict(Counter(record["supplier"] for record in records)), "images": image_stats, "publicBytes": args.public_output.stat().st_size, "privateOutput": str(args.private_output), "publicOutput": str(args.public_output)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
